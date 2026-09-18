import requests
from bs4 import BeautifulSoup
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# 크롤링 대상 URL
TARGET_URL = (
    "https://search.naver.com/search.naver?"
    "where=nexearch&sm=top_sug.pre&fbm=0&acr=1&acq=%EB%B0%98%EB%8F%84%EC%B2%B4"
    "&qdt=0&ie=utf8&query=%EB%B0%98%EB%8F%84%EC%B2%B4&ackey=oerk11xq"
)

# 네이버 서버의 봇 차단을 방지하기 위한 User-Agent 헤더
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    )
}


def crawl_naver_news_search(search_url):
    """
    네이버 통합검색 결과 페이지에서 뉴스 목록(제목, 언론사, 날짜, 요약, 링크)을 추출합니다.
    """
    response = requests.get(search_url, headers=HEADERS)
    response.raise_for_status()

    soup = BeautifulSoup(response.text, "html.parser")

    # 뉴스 섹션 내 기사 목록 컨테이너 탐색
    news_list = soup.select_one(".fds-news-item-list-desk")
    if not news_list:
        # 혹시 모를 대체 레이아웃 대비
        news_h2 = soup.find(lambda e: e.name in ["h2", "h3"] and e.get_text(strip=True) == "뉴스")
        news_list = news_h2.find_parent("div", class_=lambda c: c and "sc_new" in c) if news_h2 else soup

    # 뉴스 리스트 안의 각 기사 카드 단위(div) 탐색
    cards = [c for c in news_list.find_all("div", recursive=False) if c.name == "div"]
    
    articles = []

    for card in cards:
        # 1. 제목 태그 (data-heatmap-target=".tit")
        tit_el = card.select_one('a[data-heatmap-target=".tit"]')
        if not tit_el:
            continue

        title = tit_el.get_text(separator=" ", strip=True).replace("새 창 열림", "").strip()
        origin_url = tit_el.get("href", "")

        # 2. 언론사 정보
        press_el = card.select_one(".sds-comps-profile-info-title")
        press = press_el.get_text(strip=True).replace("새 창 열림", "").strip() if press_el else "미상"

        # 3. 작성 / 발행 시점
        date_el = card.select_one(".sds-comps-profile-info-subtext")
        date_text = date_el.get_text(strip=True) if date_el else ""

        # 4. 기사 본문 요약 (data-heatmap-target=".body")
        body_el = card.select_one('a[data-heatmap-target=".body"]')
        snippet = body_el.get_text(strip=True).replace("새 창 열림", "").strip() if body_el else ""

        # 5. 네이버뉴스 바로가기 링크 (data-heatmap-target=".nav")
        nav_el = card.select_one('a[data-heatmap-target=".nav"]')
        naver_url = nav_el.get("href", "") if nav_el else ""

        articles.append({
            "press": press,
            "date": date_text,
            "title": title,
            "snippet": snippet,
            "origin_url": origin_url,
            "naver_url": naver_url,
        })

    return articles


def crawl_naver_news_article_body(naver_news_url):
    """
    네이버 뉴스 링크(n.news.naver.com)로 접속하여 기사의 '전체 본문 원문'을 크롤링합니다.
    """
    if not naver_news_url or "n.news.naver.com" not in naver_news_url:
        return ""

    try:
        response = requests.get(naver_news_url, headers=HEADERS, timeout=5)
        response.raise_for_status()

        soup = BeautifulSoup(response.text, "html.parser")

        # 네이버 뉴스 기사 본문 영역 셀렉터
        body_el = soup.select_one("#newsct_article") or soup.select_one("#dic_area") or soup.select_one("#articeBody")
        if not body_el:
            return ""

        # 본문 내 불필요한 요소 제거 (기자 정보, 이미지 캡션 등)
        for tag in body_el.select(".byline, .reporter_area, em.img_desc, script, style"):
            tag.decompose()

        return body_el.get_text(separator="\n", strip=True)
    except Exception as e:
        return f"[본문 수집 오류: {e}]"


def save_to_excel(articles, filename="naver_result.xlsx"):
    """
    크롤링한 기사 목록을 openpyxl을 사용하여 엑셀 파일로 저장합니다.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "네이버뉴스_검색결과"

    # 헤더 정의
    headers = ["번호", "언론사", "작성일시", "기사 제목", "요약 내용", "원문 링크", "네이버뉴스 링크"]
    ws.append(headers)

    # 데이터 행 추가
    for idx, art in enumerate(articles, 1):
        ws.append([
            idx,
            art.get("press", ""),
            art.get("date", ""),
            art.get("title", ""),
            art.get("snippet", ""),
            art.get("origin_url", ""),
            art.get("naver_url", "")
        ])

    # 서식 및 스타일 설정
    header_fill = PatternFill(start_color="03C75A", end_color="03C75A", fill_type="solid")  # 네이버 시그니처 그린
    header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="맑은 고딕", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0")
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    # 1. 헤더 스타일 적용
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    # 2. 데이터 행 스타일 적용
    for row_idx in range(2, len(articles) + 2):
        ws.row_dimensions[row_idx].height = 22
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [1, 2, 3]:  # 번호, 언론사, 작성일시
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # 3. 열 너비 지정
    col_widths = {
        "A": 8,   # 번호
        "B": 16,  # 언론사
        "C": 14,  # 작성일시
        "D": 45,  # 기사 제목
        "E": 55,  # 요약 내용
        "F": 35,  # 원문 링크
        "G": 35   # 네이버뉴스 링크
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(filename)
    print(f"\n[저장 완료] 엑셀 파일이 성공적으로 생성되었습니다 -> {filename}")


if __name__ == "__main__":
    print("=== [1] 네이버 검색 결과 기사 목록 크롤링 ===")
    articles = crawl_naver_news_search(TARGET_URL)
    print(f"총 {len(articles)}개의 기사를 찾았습니다.\n")

    for idx, art in enumerate(articles, 1):
        print(f"[{idx}] {art['title']}")
        print(f"    - 언론사: {art['press']} ({art['date']})")
        print(f"    - 요약문: {art['snippet']}")
        print(f"    - 원문 링크: {art['origin_url']}")
        print(f"    - 네이버뉴스: {art['naver_url']}")
        print("-" * 70)

    # 엑셀 파일로 저장
    save_to_excel(articles, "naver_result.xlsx")

    # 첫 번째 기사의 네이버 뉴스 본문 전체 크롤링 예시
    if articles and articles[0]["naver_url"]:
        first_article = articles[0]
        print("\n=== [2] 네이버 뉴스 상세 페이지 본문 전체 크롤링 (1번 기사 예시) ===")
        print(f"대상 기사: {first_article['title']}")
        print(f"대상 URL: {first_article['naver_url']}\n")

        full_content = crawl_naver_news_article_body(first_article["naver_url"])
        print("[기사 본문 원문 내용 (앞부분 300자)]:")
        print(full_content[:300] + "\n...")
