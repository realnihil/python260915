import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

# 네이버 증권 데이터 요청 헤더
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    )
}


def fetch_naver_stocks(market="KOSPI", top_n=50):
    """
    네이버 증권 내부 API를 호출하여 시가총액 순으로 주식 목록을 수집합니다.

    Parameters:
        market (str): 'KOSPI' 또는 'KOSDAQ'
        top_n (int): 수집할 상위 종목 수 (기본값: 50)

    Returns:
        list[dict]: 수집된 주식 정보 딕셔너리 리스트
    """
    page_size = min(top_n, 100)  # 한 페이지당 최대 100개 요청 가능
    total_pages = (top_n + page_size - 1) // page_size

    stocks_data = []

    print(f"[*] 네이버 증권 [{market}] 시가총액 상위 {top_n}개 종목 데이터 수집 중...")

    rank = 1
    for page in range(1, total_pages + 1):
        url = f"https://m.stock.naver.com/api/stocks/marketValue/{market}?page={page}&pageSize={page_size}"
        
        try:
            response = requests.get(url, headers=HEADERS, timeout=10)
            response.raise_for_status()
            data = response.json()
        except requests.exceptions.RequestException as e:
            print(f"[!] API 요청 실패 (페이지 {page}): {e}")
            break

        stocks = data.get("stocks", [])
        if not stocks:
            break

        for item in stocks:
            if rank > top_n:
                break

            # 가격 변동 부호 및 기호 처리
            diff_type = item.get("compareToPreviousPrice", {}).get("name", "")
            raw_diff = item.get("compareToPreviousClosePrice", "0")
            raw_ratio = item.get("fluctuationsRatio", "0.0")

            if diff_type == "RISING":
                diff = f"+{raw_diff}"
                ratio = f"+{raw_ratio}%"
            elif diff_type == "FALLING":
                diff = raw_diff if raw_diff.startswith("-") else f"-{raw_diff}"
                ratio = f"{raw_ratio}%" if raw_ratio.startswith("-") else f"-{raw_ratio}%"
            else:
                diff = "0"
                ratio = "0.00%"

            stocks_data.append({
                "순위": rank,
                "종목코드": item.get("itemCode", ""),
                "종목명": item.get("stockName", ""),
                "시장": market,
                "현재가(원)": item.get("closePrice", "0"),
                "전일대비(원)": diff,
                "등락률": ratio,
                "변동구분": diff_type,  # 서식용 (RISING, FALLING, STEADY)
                "거래량(주)": item.get("accumulatedTradingVolume", "0"),
                "거래대금": item.get("accumulatedTradingValueKrwHangeul", "0"),
                "시가총액": item.get("marketValueHangeul", ""),
                "상세링크": item.get("newPcUrl", f"https://stock.naver.com/domestic/stock/{item.get('itemCode', '')}")
            })
            rank += 1

    print(f"[+] 총 {len(stocks_data)}개 종목 수집 완료!")
    return stocks_data


def save_to_excel(stocks, filename="naver_stock_ranking.xlsx"):
    """
    수집한 주식 데이터를 스타일이 적용된 Excel 파일로 저장합니다.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "시가총액_순위"

    # 시트 눈금선 표시
    ws.views.sheetView[0].showGridLines = True

    # 1. 메인 타이틀 행
    title_text = f"네이버 증권 주식 시세 데이터 ({datetime.now().strftime('%Y-%m-%d %H:%M')})"
    ws.merge_cells("A1:K1")
    ws["A1"] = title_text
    ws["A1"].font = Font(name="맑은 고딕", size=14, bold=True, color="1E293B")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # 2. 헤더 행 정의
    headers = [
        "순위", "종목코드", "종목명", "시장",
        "현재가(원)", "전일대비(원)", "등락률",
        "거래량(주)", "거래대금", "시가총액", "상세정보"
    ]
    ws.append([])  # 2행: 여백
    ws.append(headers)  # 3행: 헤더
    ws.row_dimensions[3].height = 28

    header_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # 네이비 블루
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # 3. 데이터 행 추가
    rise_font = Font(name="맑은 고딕", size=10, color="DC2626")    # 빨간색 (상승)
    fall_font = Font(name="맑은 고딕", size=10, color="2563EB")    # 파란색 (하락)
    default_font = Font(name="맑은 고딕", size=10, color="334155")
    even_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for i, stock in enumerate(stocks, start=4):
        ws.row_dimensions[i].height = 22
        
        row_data = [
            stock["순위"],
            stock["종목코드"],
            stock["종목명"],
            stock["시장"],
            stock["현재가(원)"],
            stock["전일대비(원)"],
            stock["등락률"],
            stock["거래량(주)"],
            stock["거래대금"],
            stock["시가총액"],
            "바로가기"
        ]
        ws.append(row_data)

        # 행 스타일 적용
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=i, column=col_idx)
            cell.border = thin_border
            cell.font = default_font

            # 짝수행 은은한 배경색
            if i % 2 == 0:
                cell.fill = even_row_fill

            # 텍스트 정렬
            if col_idx in [1, 2, 4, 11]:  # 순위, 코드, 시장, 링크
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [3]:  # 종목명
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:  # 가격/수치 정보
                cell.alignment = Alignment(horizontal="right", vertical="center")

            # 전일대비 및 등락률 텍스트 색상 (상승: 빨강, 하락: 파랑)
            if col_idx in [6, 7]:
                if stock["변동구분"] == "RISING":
                    cell.font = rise_font
                elif stock["변동구분"] == "FALLING":
                    cell.font = fall_font

            # 링크 열 하이퍼링크 처리
            if col_idx == 11:
                cell.hyperlink = stock["상세링크"]
                cell.font = Font(name="맑은 고딕", size=10, color="2563EB", underline="single")

    # 4. 열 너비 자동 조정
    for col in ws.columns:
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in [1, 2]:
                continue
            if cell.value:
                # 한글 고려 대략적인 길이 계산
                text = str(cell.value)
                length = sum(2 if ord(c) > 127 else 1 for c in text)
                max_len = max(max_len, length)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

    wb.save(filename)
    print(f"[+] 엑셀 파일 저장 완료: {filename}")


if __name__ == "__main__":
    # 코스피 시가총액 상위 50개 종목 수집
    kospi_stocks = fetch_naver_stocks(market="KOSPI", top_n=50)

    # 수집 결과 콘솔 미리보기 (상위 5개)
    print("\n[미리보기] 상위 5개 종목:")
    print(f"{'순위':<4} {'종목명':<15} {'현재가':<10} {'등락률':<10} {'시가총액':<15}")
    print("-" * 60)
    for s in kospi_stocks[:5]:
        print(f"{s['순위']:<4} {s['종목명']:<15} {s['현재가(원)']:<10} {s['등락률']:<10} {s['시가총액']:<15}")

    # 엑셀 파일 저장
    save_to_excel(kospi_stocks, filename="naver_kospi_top50.xlsx")
