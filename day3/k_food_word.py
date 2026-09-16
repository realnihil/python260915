from datetime import datetime
from html import unescape
from pathlib import Path
import re
from urllib.parse import quote
from urllib.request import Request, urlopen
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


SEARCH_TERMS = ["K-food", "한식", "한국 음식", "K-푸드"]
ARTICLE_LIMIT = 5
OUTPUT_PATH = Path(__file__).resolve().parent.parent / "k-food.docx"


def clean_text(value):
    text = unescape(value or "")
    return re.sub(r"<[^>]+>", "", text).strip()


def collect_articles(search_term, limit=ARTICLE_LIMIT):
    query = quote(f"{search_term} when:30d")
    rss_url = f"https://news.google.com/rss/search?q={query}&hl=ko&gl=KR&ceid=KR:ko"
    request = Request(rss_url, headers={"User-Agent": "Mozilla/5.0"})

    with urlopen(request, timeout=15) as response:
        root = ET.fromstring(response.read())

    articles = []
    for item in root.findall("./channel/item")[:limit]:
        articles.append(
            {
                "title": clean_text(item.findtext("title")),
                "description": clean_text(item.findtext("description")),
                "link": item.findtext("link", "").strip(),
                "published": item.findtext("pubDate", "").strip(),
                "source": clean_text(item.findtext("source")),
            }
        )
    return articles


def build_document(articles_by_term):
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Malgun Gothic"
    normal_style.font.size = Pt(10)

    title = document.add_heading("K-food 자료 수집 보고서", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        f"작성일: {datetime.now():%Y-%m-%d %H:%M}\n"
        "최근 30일간 Google News RSS에서 수집한 K-food 관련 자료입니다."
    )

    for search_term, articles in articles_by_term.items():
        document.add_heading(f"검색어: {search_term}", level=1)
        if not articles:
            document.add_paragraph("수집된 자료가 없습니다.")
            continue

        for index, article in enumerate(articles, start=1):
            document.add_heading(f"{index}. {article['title']}", level=2)
            document.add_paragraph(f"출처: {article['source'] or '알 수 없음'}")
            document.add_paragraph(f"발행일: {article['published'] or '알 수 없음'}")
            if article["description"]:
                document.add_paragraph(article["description"])
            document.add_paragraph(f"원문 링크: {article['link']}")

    document.save(OUTPUT_PATH)


def main():
    articles_by_term = {}
    for search_term in SEARCH_TERMS:
        try:
            articles_by_term[search_term] = collect_articles(search_term)
        except Exception as error:
            print(f"'{search_term}' 수집 실패: {error}")
            articles_by_term[search_term] = []

    build_document(articles_by_term)
    total = sum(len(articles) for articles in articles_by_term.values())
    print(f"{total}개의 자료를 {OUTPUT_PATH}에 저장했습니다.")


if __name__ == "__main__":
    main()